"""
國泰帳戶「6/8 起報酬」結算（2026-10-09 起）：讀 Wealth OS V4.4 的 03_國泰漲跌 每日表，
仿「7/3 起加碼損益」報告格式產出 Markdown＋HTML。

資料口徑（直接沿用 03_國泰漲跌，不另抓股價）：
  總市值 MV、持有成本（＝MV − 未實現）、淨投入 CF（＝當日成本變動）、單日損益（真實）＝ΔMV − CF
算法：
  淨損益     ＝ 期末 MV − 基準日 MV − ΣCF（＝期間單日損益加總，程式會交叉驗算）
  簡單報酬率 ＝ 淨損益 ÷（基準日 MV＋ΣCF）
  Modified Dietz ＝ 淨損益 ÷（基準日 MV＋Σ CF×剩餘日數權重）
  時間加權 TWR ＝ Π(1＋r) − 1，r ＝ 單日損益 ÷（前日 MV＋當日 CF）——假設投入當天開盤就買進
基準日＝起算日前最後一個交易日（6/8 起算 → 6/5 收盤）。

用法：
  python3 cathay_since_pnl.py                         # 從雲端下載 xlsm，起算 2026-06-08
  python3 cathay_since_pnl.py --xlsm 本機.xlsm --since 2026-06-08
"""
import argparse
import datetime as dt
import html
import json
import os
import sys
import tempfile

import openpyxl

SHEET = "03_國泰漲跌"
HIST_SHEET = "21_國泰資產歷史"
OUT_DIR = os.path.dirname(os.path.abspath(__file__))
WD = "一二三四五六日"
# 含帳戶金額的個人註記另存、列入 .gitignore（jimmy-agent 是公開 repo）；沒有這個檔就略過那幾段
NOTES_FILE = os.path.join(OUT_DIR, "cathay_report_notes.json")


def load_notes():
    try:
        with open(NOTES_FILE, encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        return {}


# ───────────────────────── 讀資料 ─────────────────────────
def load_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows, seen = [], set()
    for r in wb[SHEET].iter_rows(values_only=True):
        d, mv, unreal, pnl, cf = (list(r) + [None] * 5)[:5]
        if isinstance(d, dt.datetime) and isinstance(mv, (int, float)) and d.date() not in seen:
            seen.add(d.date())
            rows.append(dict(d=d.date(), mv=mv, cost=mv - unreal, pnl=pnl or 0, cf=cf or 0))
    rows.sort(key=lambda x: x["d"])
    # 交叉驗證 21_國泰資產歷史（03 由它產生，兩邊應一致）
    hist = {}
    for r in wb[HIST_SHEET].iter_rows(min_row=2, values_only=True):
        if isinstance(r[0], dt.datetime):
            hist[r[0].date()] = (r[1], r[2])
    diff = [x["d"] for x in rows if x["d"] in hist and
            (round(hist[x["d"]][0]) != round(x["mv"]) or round(hist[x["d"]][1]) != round(x["cost"]))]
    return rows, diff, len(hist)


def fetch_xlsm():
    import update_wealth_os as w
    _, meta, path, _ = w.download_xlsm(w.get_creds(), tempfile.mkdtemp())
    return path, meta["modifiedTime"]


# ───────────────────────── 計算 ─────────────────────────
def twr(days, prev_mv):
    g = 1.0
    for x in days:
        base = prev_mv + x["cf"]
        g *= 1 + (x["pnl"] / base if base else 0)
        prev_mv = x["mv"]
    return g - 1


def compute(rows, since):
    i0 = next(i for i, x in enumerate(rows) if x["d"] >= since)
    base, period = rows[i0 - 1], rows[i0:]
    end = period[-1]
    cf = sum(x["cf"] for x in period)
    pnl = end["mv"] - base["mv"] - cf
    assert abs(pnl - sum(x["pnl"] for x in period)) < 1, "ΔMV−ΣCF 與單日損益加總不符"
    span = (end["d"] - base["d"]).days
    dietz_den = base["mv"] + sum(x["cf"] * (end["d"] - x["d"]).days / span for x in period)

    months, prev = [], base
    for ym in sorted({(x["d"].year, x["d"].month) for x in period}):
        ds = [x for x in period if (x["d"].year, x["d"].month) == ym]
        months.append(dict(ym=ym, n=len(ds), first=ds[0]["d"], start=prev["mv"], cf=sum(x["cf"] for x in ds),
                           end=ds[-1]["mv"], pnl=sum(x["pnl"] for x in ds), twr=twr(ds, prev["mv"]),
                           up=sum(x["pnl"] > 0 for x in ds)))
        prev = ds[-1]

    cum, peak, mdd, curve = 0, 0, 0, []
    prev_mv = base["mv"]
    for x in period:
        x["r"] = x["pnl"] / (prev_mv + x["cf"])
        prev_mv = x["mv"]
        cum += x["pnl"]
        peak = max(peak, cum)
        mdd = min(mdd, cum - peak)
        curve.append((x["d"], cum))
    return dict(base=base, end=end, period=period, cf=cf, pnl=pnl,
                simple=pnl / (base["mv"] + cf), dietz=pnl / dietz_den, twr=twr(period, base["mv"]),
                months=months, curve=curve, mdd=mdd, peak=max(c for _, c in curve),
                up=sum(x["pnl"] > 0 for x in period), down=sum(x["pnl"] < 0 for x in period),
                flows=[x for x in period if x["cf"]],
                best=sorted(period, key=lambda x: -x["pnl"])[:5],
                worst=sorted(period, key=lambda x: x["pnl"])[:5])


# ───────────────────────── 輸出共用 ─────────────────────────
md_ = lambda d: f"{d.month}/{d.day:02d}"
ymd = lambda d: f"{d:%Y/%m/%d}"
f0 = lambda v: f"{v:,.0f}"
sg = lambda v: f"{v:+,.0f}" if round(v) else "0"
pc = lambda v: f"{v * 100:+.2f}%"
wan = lambda v: f"{v / 10000:.0f}" if v >= 100000 else f"{v / 10000:.1f}"


def mlabel(m):
    y, mo = m["ym"]
    return f"{mo}月" + (f"（{md_(m['first'])} 起）" if m["first"].day > 7 else "")


def title(R):
    verb = "賺了" if R["pnl"] >= 0 else "賠了"
    return f"6/8 起國泰帳戶淨投入 {wan(R['cf'])} 萬，{verb} {wan(abs(R['pnl']))} 萬"


# ───────────────────────── Markdown ─────────────────────────
def build_md(R, since, src_time):
    b, e = R["base"], R["end"]
    N = load_notes()
    L = [f"Wealth OS V4.4 · {SHEET} · 結算日 {ymd(e['d'])} 收盤", "",
         f"# {title(R)}", "",
         f"以 {ymd(b['d'])} 收盤市值為基準，計入 {ymd(since)}（含）起 {len(R['period'])} 個交易日的淨投入與每日損益，"
         f"結算到 {ymd(e['d'])} 收盤。", "",
         f"淨損益 {sg(R['pnl'])}　簡單報酬 {pc(R['simple'])}　時間加權 {pc(R['twr'])}", "",
         "| 項目 | 金額 |", "| --- | ---: |",
         f"| 基準市值（{md_(b['d'])} 收盤） | {f0(b['mv'])} |",
         f"| 期間淨投入（持有成本增加） | {sg(R['cf'])} |",
         f"| 期末市值（{md_(e['d'])} 收盤） | {f0(e['mv'])} |",
         f"| **淨損益（期末 − 基準 − 淨投入）** | **{sg(R['pnl'])}** |", "",
         "## 報酬率（三種口徑）", "",
         "| 口徑 | 報酬率 | 說明 |", "| --- | ---: | --- |",
         f"| 簡單報酬 | {pc(R['simple'])} | 淨損益 ÷（基準市值＋淨投入）＝把所有錢都當期初就投入 |",
         f"| Modified Dietz | {pc(R['dietz'])} | 依資金實際在場天數加權，最接近「這筆錢的報酬」 |",
         f"| 時間加權 TWR | {pc(R['twr'])} | 逐日連乘，排除投入時點影響，適合跟大盤／ETF 比 |", "",
         "## 各月貢獻", "",
         "| 月份 | 交易日 | 月初市值 | 淨投入 | 月底市值 | 淨損益 | 月報酬（TWR） | 上漲天數 |",
         "| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |"]
    for m in R["months"]:
        L.append(f"| {mlabel(m)} | {m['n']} | {f0(m['start'])} | {sg(m['cf'])} | {f0(m['end'])} | "
                 f"{sg(m['pnl'])} | {pc(m['twr'])} | {m['up']}/{m['n']} |")
    L += [f"| **合計** | {len(R['period'])} | {f0(b['mv'])} | {sg(R['cf'])} | {f0(e['mv'])} | "
          f"**{sg(R['pnl'])}** | {pc(R['twr'])} | {R['up']}/{len(R['period'])} |", "",
          "## 期間資金投入", "", "淨投入＝當日持有成本變動（買進為正）", "",
          "| 日期 | 淨投入 | 當日市值 | 當日損益 |", "| --- | ---: | ---: | ---: |"]
    L += [f"| {ymd(x['d'])} | {sg(x['cf'])} | {f0(x['mv'])} | {sg(x['pnl'])} |" for x in R["flows"]]
    L += [f"| **合計** | **{sg(R['cf'])}** | | |", "",
          "## 單日極值", "", "| 最佳 5 日 | 損益 | 報酬 | 最差 5 日 | 損益 | 報酬 |",
          "| --- | ---: | ---: | --- | ---: | ---: |"]
    for a, z in zip(R["best"], R["worst"]):
        L.append(f"| {ymd(a['d'])}（{WD[a['d'].weekday()]}） | {sg(a['pnl'])} | {pc(a['r'])} | "
                 f"{ymd(z['d'])}（{WD[z['d'].weekday()]}） | {sg(z['pnl'])} | {pc(z['r'])} |")
    L += ["", f"- 上漲 {R['up']} 天／下跌 {R['down']} 天；累積損益最高 {sg(R['peak'])}，"
              f"期間最大回落 {sg(R['mdd'])}（由高點回吐的金額）。", "",
          "## 計算依據與限制", "",
          "### 資料口徑",
          f"- 直接取 `{SHEET}` 每日表：總市值、未實現損益、單日損益（真實）、淨投入資金；"
          f"已與 `{HIST_SHEET}` 逐日比對市值與持有成本。",
          "- 單日損益（真實）＝ 當日市值變動 − 淨投入，所以買進不會被算成賺錢。",
          "- 淨損益已用「期末 − 基準 − 淨投入」與「單日損益加總」雙向驗算，兩者一致。", "",
          "### 報酬率",
          "- 日報酬分母＝前日市值＋當日淨投入（假設投入當天開盤就買進）；"
          "表上「單日報酬率」欄分母只用前日市值，投入當天會被放大。",
          "- 未動用的現金不計入分母，只計實際買成股票的部分。"] + N.get("md_rate", []) + ["",
          "### 限制",
          "- 未含尚未賣出的手續費與證交稅（7/3 報告有扣，這份是帳戶市值口徑、沒有逐檔持股）。",
          "- 現金股利若未反映在市值或成本，不會出現在損益中。"] + N.get("md_extra", []) + ["",
          f"產出日期 {dt.date.today():%Y/%m/%d}。資料來源：雲端 Jimmy_Wealth_OS_Master_V4.4_Google.xlsm"
          f"（{src_time}）。"]
    return "\n".join(L) + "\n"


# ───────────────────────── HTML ─────────────────────────
CSS = open(os.path.join(OUT_DIR, "investment_since_0703_report.html"), encoding="utf-8").read()
CSS = CSS[CSS.index("<style>"):CSS.index("</style>") + 8]
CSS = CSS.replace("</style>", """
.kpis{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:12px}
.kpi{background:var(--surface);border:1px solid var(--line);border-radius:10px;padding:14px 16px;display:grid;gap:4px}
.kpi .k{font-size:12px;color:var(--muted)} .kpi .v{font-family:var(--f-num);font-size:22px;font-weight:600}
.kpi .d{font-size:12px;color:var(--muted)}
.chart{background:var(--surface);border:1px solid var(--line);border-radius:10px;padding:12px 14px}
.chart svg{display:block;width:100%;height:auto}
.chart text{font-family:var(--f-num);font-size:11px;fill:var(--muted)}
@media (max-width:640px){.kpis{grid-template-columns:1fr}}
</style>""")


def cls(v):
    return "up" if v > 0 else "down" if v < 0 else ""


def svg_curve(curve, flows):
    W, H, PL, PR, PT, PB = 900, 260, 64, 12, 14, 26
    vals = [v for _, v in curve] + [0]
    lo, hi = min(vals), max(vals)
    pad = (hi - lo) * .08 or 1
    lo, hi = lo - pad, hi + pad
    n = len(curve)
    X = lambda i: PL + (W - PL - PR) * i / max(n - 1, 1)
    Y = lambda v: PT + (H - PT - PB) * (hi - v) / (hi - lo)
    pts = " ".join(f"{X(i):.1f},{Y(v):.1f}" for i, (_, v) in enumerate(curve))
    area = f"{X(0):.1f},{Y(0):.1f} {pts} {X(n - 1):.1f},{Y(0):.1f}"
    step = 10 ** max(len(str(int(hi - lo))) - 1, 0)
    step = step * (5 if (hi - lo) / step > 8 else 2 if (hi - lo) / step > 4 else 1)
    grid = []
    t = (lo // step + 1) * step
    while t < hi:
        stroke = 'stroke="var(--muted)" stroke-width="1.5"' if t == 0 else 'stroke="var(--line)"'
        grid.append(f'<line x1="{PL}" x2="{W - PR}" y1="{Y(t):.1f}" y2="{Y(t):.1f}" {stroke}/>'
                    f'<text x="{PL - 6}" y="{Y(t) + 4:.1f}" text-anchor="end">{t / 10000:+.0f}萬</text>')
        t += step
    xl, lastm = [], None
    for i, (d, _) in enumerate(curve):
        if d.month != lastm:
            lastm = d.month
            xl.append(f'<text x="{X(i):.1f}" y="{H - 8}" text-anchor="start">{d.month}/{d.day}</text>')
    fd = {x["d"] for x in flows}
    dots = "".join(f'<circle cx="{X(i):.1f}" cy="{Y(v):.1f}" r="3.2" fill="var(--accent)"><title>{md_(d)} 投入</title></circle>'
                   for i, (d, v) in enumerate(curve) if d in fd)
    end_d, end_v = curve[-1]
    return (f'<svg viewBox="0 0 {W} {H}" role="img" aria-label="累積淨損益走勢">{"".join(grid)}{"".join(xl)}'
            f'<polygon points="{area}" fill="var(--up-soft)" opacity=".7"/>'
            f'<polyline points="{pts}" fill="none" stroke="var(--up)" stroke-width="2"/>{dots}'
            f'<circle cx="{X(n - 1):.1f}" cy="{Y(end_v):.1f}" r="4" fill="var(--up)"/>'
            f'<text x="{X(n - 1) - 6:.1f}" y="{Y(end_v) - 9:.1f}" text-anchor="end" style="fill:var(--ink);font-weight:600">'
            f'{sg(end_v)}</text></svg>')


def bars(months):
    mx = max(abs(m["pnl"]) for m in months) or 1
    lo = min(0, min(m["pnl"] for m in months)) / mx
    zero = -lo / (1 - lo) * 100
    out = []
    for m in months:
        w = abs(m["pnl"]) / mx / (1 - lo) * 100
        left = zero if m["pnl"] >= 0 else zero - w
        out.append(f'<div class="bar-row"><span class="name">{html.escape(mlabel(m).split("（")[0])}</span>'
                   f'<div class="track" style="--zero:{zero:.1f}%"><i class="{"pos" if m["pnl"] >= 0 else "neg"}" '
                   f'style="left:{left:.1f}%;width:{w:.1f}%"></i></div>'
                   f'<span class="amt {cls(m["pnl"])}">{sg(m["pnl"])}</span></div>')
    return "".join(out)


def build_html(R, since, src_time):
    b, e = R["base"], R["end"]
    N = load_notes()
    rate_li = "".join(f"<li>{t}</li>" for t in N.get("html_rate", []))
    cards = "".join(N.get("html_cards", []))
    td = lambda v, f=sg: f'<td class="{cls(v)}">{f(v)}</td>'
    mrows = "".join(
        f'<tr><td>{html.escape(mlabel(m))}</td><td>{m["n"]}</td><td>{f0(m["start"])}</td><td>{sg(m["cf"])}</td>'
        f'<td>{f0(m["end"])}</td>{td(m["pnl"])}{td(m["twr"], pc)}<td>{m["up"]}/{m["n"]}</td></tr>'
        for m in R["months"])
    frows = "".join(f'<tr><td>{ymd(x["d"])}</td><td>{sg(x["cf"])}</td><td>{f0(x["mv"])}</td>{td(x["pnl"])}</tr>'
                    for x in R["flows"])
    xrows = "".join(
        f'<tr><td>{ymd(a["d"])}（{WD[a["d"].weekday()]}）</td>{td(a["pnl"])}{td(a["r"], pc)}'
        f'<td>{ymd(z["d"])}（{WD[z["d"].weekday()]}）</td>{td(z["pnl"])}{td(z["r"], pc)}</tr>'
        for a, z in zip(R["best"], R["worst"]))
    pcls = cls(R["pnl"])
    return f"""<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>國泰 6/8 起報酬</title>
<link rel="stylesheet" href="https://fonts.googleapis.com/css2?family=Noto+Sans+TC:wght@400;500;700&family=Noto+Serif+TC:wght@600;700&family=IBM+Plex+Mono:wght@400;500;600&display=swap">
{CSS}</head><body>
<div class="wrap">
<header>
  <div class="eyebrow">Wealth OS V4.4 · {SHEET} · 結算日 {ymd(e["d"])} 收盤</div>
  <h1>{html.escape(title(R))}</h1>
  <p class="lede">以 {ymd(b["d"])} 收盤市值為基準，計入 {ymd(since)}（含）起 {len(R["period"])} 個交易日的淨投入與每日損益，結算到 {ymd(e["d"])} 收盤。</p>
</header>

<section class="hero">
  <div class="big"><span class="label">淨損益（扣除期間投入）</span>
    <span class="value {pcls}" style="color:var(--{"up" if R["pnl"] >= 0 else "down"})">{sg(R["pnl"])}</span>
    <span class="pct" style="color:var(--{"up" if R["pnl"] >= 0 else "down"})">簡單 {pc(R["simple"])}　·　TWR {pc(R["twr"])}</span>
    <span class="label">基準 {f0(b["mv"])} ＋ 投入 {f0(R["cf"])} → 期末 {f0(e["mv"])}</span></div>
  <div class="ledger">
    <div><span>基準市值（{md_(b["d"])} 收盤）</span><span>{f0(b["mv"])}</span></div>
    <div><span>期間淨投入</span><span>{sg(R["cf"])}</span></div>
    <div><span>期末市值（{md_(e["d"])} 收盤）</span><span>{f0(e["mv"])}</span></div>
    <div><span>上漲／下跌天數</span><span>{R["up"]} ／ {R["down"]}</span></div>
    <div><span>淨損益</span><span class="{pcls}">{sg(R["pnl"])}</span></div>
  </div>
</section>

<section>
  <div class="sec-head"><h2>報酬率（三種口徑）</h2><p>投入時點不同，口徑就不同</p></div>
  <div class="kpis">
    <div class="kpi"><span class="k">簡單報酬</span><span class="v {cls(R["simple"])}">{pc(R["simple"])}</span><span class="d">淨損益 ÷（基準市值＋淨投入）</span></div>
    <div class="kpi"><span class="k">Modified Dietz</span><span class="v {cls(R["dietz"])}">{pc(R["dietz"])}</span><span class="d">依資金在場天數加權，最貼近「這筆錢的報酬」</span></div>
    <div class="kpi"><span class="k">時間加權 TWR</span><span class="v {cls(R["twr"])}">{pc(R["twr"])}</span><span class="d">逐日連乘、排除投入時點，可直接跟大盤比</span></div>
  </div>
</section>

<section>
  <div class="sec-head"><h2>累積淨損益走勢</h2><p>藍點＝有資金投入的日子；最高 {sg(R["peak"])}、最大回落 {sg(R["mdd"])}</p></div>
  <div class="chart">{svg_curve(R["curve"], R["flows"])}</div>
</section>

<section>
  <div class="sec-head"><h2>各月貢獻</h2><p>紅色賺、綠色賠</p></div>
  <div class="bars">{bars(R["months"])}</div>
</section>

<section>
  <div class="sec-head"><h2>逐月明細</h2><p>表格可左右捲動</p></div>
  <div class="table-box"><table>
    <thead><tr><th>月份</th><th>交易日</th><th>月初市值</th><th>淨投入</th><th>月底市值</th><th>淨損益</th><th>月報酬（TWR）</th><th>上漲天數</th></tr></thead>
    <tbody>{mrows}</tbody>
    <tfoot><tr><td>合計</td><td>{len(R["period"])}</td><td>{f0(b["mv"])}</td><td>{sg(R["cf"])}</td><td>{f0(e["mv"])}</td>{td(R["pnl"])}{td(R["twr"], pc)}<td>{R["up"]}/{len(R["period"])}</td></tr></tfoot>
  </table></div>
</section>

<section>
  <div class="sec-head"><h2>期間資金投入</h2><p>淨投入＝當日持有成本變動（買進為正）</p></div>
  <div class="table-box"><table style="min-width:520px">
    <thead><tr><th>日期</th><th>淨投入</th><th>當日市值</th><th>當日損益</th></tr></thead>
    <tbody>{frows}</tbody>
    <tfoot><tr><td>合計</td><td>{sg(R["cf"])}</td><td></td><td></td></tr></tfoot>
  </table></div>
</section>

<section>
  <div class="sec-head"><h2>單日極值</h2><p>日報酬分母＝前日市值＋當日投入</p></div>
  <div class="table-box"><table style="min-width:640px">
    <thead><tr><th>最佳 5 日</th><th>損益</th><th>報酬</th><th>最差 5 日</th><th>損益</th><th>報酬</th></tr></thead>
    <tbody>{xrows}</tbody>
  </table></div>
</section>

<section>
  <h2>計算依據與限制</h2>
  <div class="notes">
    <div class="note"><h3>資料口徑</h3><ul>
      <li>取 {SHEET} 每日表的總市值、未實現損益、單日損益（真實）、淨投入資金，並與 {HIST_SHEET} 逐日比對。</li>
      <li>單日損益＝當日市值變動 − 淨投入，買進不會被算成賺錢。</li>
      <li>淨損益以「期末 − 基準 − 淨投入」與「單日損益加總」雙向驗算，一致。</li></ul></div>
    <div class="note"><h3>報酬率</h3><ul>
      <li>假設投入當天開盤就買進；表上「單日報酬率」分母只用前日市值，投入當天會被放大。</li>
      <li>未動用的現金不進分母，只計實際買成股票的部分。</li>{rate_li}</ul></div>
    <div class="note"><h3>限制</h3><ul>
      <li>帳戶市值口徑，未扣假設賣出的手續費與證交稅（7/3 報告有扣）。</li>
      <li>現金股利若沒反映在市值或成本，不會出現在損益中。</li></ul></div>
    {cards}
  </div>
</section>

<footer>產出日期 {dt.date.today():%Y/%m/%d}。資料來源：雲端 Jimmy_Wealth_OS_Master_V4.4_Google.xlsm（{src_time}）。</footer>
</div></body></html>
"""


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--since", default="2026-06-08")
    ap.add_argument("--xlsm")
    ap.add_argument("--out", default="cathay_since_0608_report")
    a = ap.parse_args()
    since = dt.date.fromisoformat(a.since)
    path, src = (a.xlsm, "本機檔") if a.xlsm else fetch_xlsm()
    rows, diff, nh = load_rows(path)
    if diff:
        print(f"⚠ {SHEET} 與 {HIST_SHEET} 不一致：{', '.join(map(str, diff))}", file=sys.stderr)
    R = compute(rows, since)
    for ext, fn in (("md", build_md), ("html", build_html)):
        p = os.path.join(OUT_DIR, f"{a.out}.{ext}")
        with open(p, "w", encoding="utf-8") as f:
            f.write(fn(R, since, src))
        print("寫入", p)
    print(f"比對 {HIST_SHEET} {nh} 列，不一致 {len(diff)} 列")
    print(f"淨損益 {sg(R['pnl'])}｜淨投入 {sg(R['cf'])}｜簡單 {pc(R['simple'])}｜Dietz {pc(R['dietz'])}｜TWR {pc(R['twr'])}")


if __name__ == "__main__":
    main()
