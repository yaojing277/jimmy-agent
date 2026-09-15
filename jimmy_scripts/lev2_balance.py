#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""lev2_balance.py — 阿良「正二人生資產負債表」計算引擎(純計算,不碰雲端)

供 update_wealth_os.py --full 重建「06_阿良資產負債表」分頁使用;也可單獨對本機 xlsm 印報表驗證:
  python3 lev2_balance.py --xlsm Jimmy_Wealth_OS_Master_V4.4_Google.xlsm

範本邏輯:
  三桶  原型(β1)=非正二非債券持股 | 正二(β2)=類別「正二ETF」 | 防守(β0)=現金+債券ETF
  金融資產總額 = 持股市值 + 現金;淨值 = 總額 − 信貸(房貸買房,不計入投資槓桿)
  本金槓桿 = 總額 ÷ 淨值;Beta% = Σ市值×β ÷ 總額;總曝險 = Σ市值×β ÷ 淨值
  生活費倍數 = 淨值 ÷ 年花費 → 查 12_設定 D:G 對照表得建議配置(原型:正二:防守 成數)
  5 年預期總報酬 = Σ 權重 × 各桶 5 年總報酬;年化 = (1+總報酬)^(1/5) − 1
  正二跌加碼梯 = 00631L 現價 × (1 − 階距×k),每階加碼固定金額,另留一筆預金

參數全部讀 12_設定(全檔唯一事實來源),缺格 raise ValueError,由呼叫端決定跳過或中止。
"""
import sys

SETTINGS_TAB = "12_設定"
LEV2_CAT = "正二ETF"
LADDER_CODE = "00631L"

# 12_設定 A/B 欄新區塊:鍵名 -> 列號
CFG_ROWS = {
    "cash": 6, "loan": 5,                                   # 既有列
    "annual_spend": 22, "step_amt": 23, "reserve": 24,
    "step_pct": 25, "steps": 26, "beta_lev": 27, "beta_base": 28,
    "r5_base": 29, "r5_lev": 30, "r5_def": 31,
}
BOND_ROW = 32                                               # 債券ETF代號(字串,逗號分隔)
TIER_FIRST_ROW, TIER_LAST_ROW = 3, 10                       # D~G:年數下限|原型|正二|防守


def _f(v, where):
    if isinstance(v, str):
        v = v.strip().lstrip("=").replace(",", "")
    try:
        return float(v)
    except (TypeError, ValueError):
        raise ValueError(f"{SETTINGS_TAB}!{where} 不是數值:{v!r}")


def read_lev2_cfg(wb):
    """讀 12_設定 的正二資產負債表參數。wb 為 openpyxl Workbook(非 data_only 亦可,需為數值)。"""
    if SETTINGS_TAB not in wb.sheetnames:
        raise ValueError(f"找不到分頁「{SETTINGS_TAB}」")
    ws = wb[SETTINGS_TAB]
    cfg = {k: _f(ws.cell(r, 2).value, f"B{r}") for k, r in CFG_ROWS.items()}
    bonds = ws.cell(BOND_ROW, 2).value
    if bonds is None or not str(bonds).strip():
        raise ValueError(f"{SETTINGS_TAB}!B{BOND_ROW}(債券ETF代號)是空的")
    cfg["bonds"] = {b.strip() for b in str(bonds).split(",") if b.strip()}
    tiers = []
    for r in range(TIER_FIRST_ROW, TIER_LAST_ROW + 1):
        yrs, b, l, d = (_f(ws.cell(r, c).value, f"{'DEFG'[c - 4]}{r}") for c in (4, 5, 6, 7))
        tiers.append({"years": yrs, "base": b, "lev": l, "def": d})
    cfg["tiers"] = sorted(tiers, key=lambda t: -t["years"])
    return cfg


def prev_month_mv(wb, today, tab="16_資產歷史"):
    """16_資產歷史 中「早於本月」的最後一筆股票市值(B 欄);找不到回 None。today 為 'YYYY/MM/DD'。"""
    if tab not in wb.sheetnames:
        return None
    y, m = (int(x) for x in str(today).split("/")[:2])
    best = None
    for a, b in wb[tab].iter_rows(min_row=2, max_col=2, values_only=True):
        try:
            ay, am = (int(x) for x in str(a).split("/")[:2])
            mv = float(b)
        except (TypeError, ValueError):
            continue
        if (ay, am) < (y, m):
            best = (str(a), mv)
    return best


def code_of(tier):
    """配置成數 → 範本代號,例 0.7/0/0.3 → '703'。"""
    return "".join(str(int(round(tier[k] * 10))) for k in ("base", "lev", "def"))


def compute(holdings, cfg, prev=None):
    """holdings: [{name, cat, sh, px}];prev: prev_month_mv() 的回傳。"""
    buckets = {"base": 0.0, "lev": 0.0, "def": 0.0}
    ladder_px = None
    for h in holdings:
        mv = h["sh"] * h["px"]
        name = str(h["name"]).strip()
        if name in cfg["bonds"]:
            buckets["def"] += mv
        elif h["cat"] == LEV2_CAT:
            buckets["lev"] += mv
        else:
            buckets["base"] += mv
        if name == LADDER_CODE:
            ladder_px = h["px"]
    stock_mv = sum(buckets.values())
    buckets["def"] += cfg["cash"]
    total = stock_mv + cfg["cash"]
    net = total - cfg["loan"]
    if total <= 0 or net <= 0:
        raise ValueError(f"金融資產總額 {total:,.0f} / 淨值 {net:,.0f} 非正數,無法計算槓桿")

    beta = {"base": cfg["beta_base"], "lev": cfg["beta_lev"], "def": 0.0}
    r5 = {"base": cfg["r5_base"], "lev": cfg["r5_lev"], "def": cfg["r5_def"]}
    exposure = sum(buckets[k] * beta[k] for k in buckets)
    actual_w = {k: buckets[k] / total for k in buckets}

    years = net / cfg["annual_spend"]
    tier = next((t for t in cfg["tiers"] if years >= t["years"]), cfg["tiers"][-1])
    target_w = {k: tier[k] for k in buckets}

    def perf(w):
        tot = sum(w[k] * r5[k] for k in w)
        return {"beta": sum(w[k] * beta[k] for k in w), "r5": tot,
                "annual": (1 + tot) ** 0.2 - 1 if tot > -1 else -1.0}

    ladder = []
    if ladder_px:
        cum = cfg["reserve"]
        for k in range(1, int(cfg["steps"]) + 1):
            drop = -cfg["step_pct"] * k
            px = ladder_px * (1 + drop)
            cum += cfg["step_amt"]
            ladder.append({"drop": drop, "px": px, "amt": cfg["step_amt"], "cum": cum,
                           "shares": int(cfg["step_amt"] / px) if px > 0 else 0})

    return {
        "buckets": buckets, "stock_mv": stock_mv, "cash": cfg["cash"], "loan": cfg["loan"],
        "total": total, "net": net, "leverage": total / net,
        "exposure": exposure, "beta_pct": exposure / total, "total_exposure": exposure / net,
        "prev": prev, "mv_diff": (stock_mv - prev[1]) if prev else None,
        "annual_spend": cfg["annual_spend"], "years": years,
        "tier": tier, "tier_code": code_of(tier), "tiers": cfg["tiers"],
        "actual_w": actual_w, "target_w": target_w,
        "rebalance": {k: target_w[k] * total - buckets[k] for k in buckets},
        "actual_perf": perf(actual_w), "target_perf": perf(target_w),
        "r5": r5, "beta": beta, "ladder_px": ladder_px, "ladder": ladder,
        "reserve": cfg["reserve"], "ladder_total": ladder[-1]["cum"] if ladder else cfg["reserve"],
    }


LABEL = {"base": "原型(β1)", "lev": "正二(β2)", "def": "防守(β0)"}


def print_report(r):
    p = print
    p("══ 正二人生資產負債表 ══")
    for k in ("base", "lev", "def"):
        p(f"  {LABEL[k]:<8} {r['buckets'][k]:>14,.0f}  {r['actual_w'][k]:6.1%}"
          f"  目標 {r['target_w'][k]:6.1%}  應增減 {r['rebalance'][k]:+14,.0f}")
    p(f"  金融資產總額 {r['total']:,.0f}|信貸 {r['loan']:,.0f}|淨值 {r['net']:,.0f}")
    p(f"  本金槓桿 {r['leverage']:.0%}|Beta {r['beta_pct']:.0%}|總曝險 {r['total_exposure']:.0%}")
    if r["prev"]:
        p(f"  股票市值 vs {r['prev'][0]}:{r['mv_diff']:+,.0f}")
    p(f"  年花費 {r['annual_spend']:,.0f} → 生活費倍數 {r['years']:.1f} 年"
      f" → 建議配置 {r['tier_code']}(≥{r['tier']['years']:g} 年)")
    a, t = r["actual_perf"], r["target_perf"]
    p(f"  實際 Beta {a['beta']:.0%} 5Y {a['r5']:.1%} 年化 {a['annual']:.1%}"
      f"|建議 Beta {t['beta']:.0%} 5Y {t['r5']:.1%} 年化 {t['annual']:.1%}")
    if r["ladder"]:
        p(f"  加碼梯(00631L 現價 {r['ladder_px']}):預金 {r['reserve']:,.0f}")
        for s in r["ladder"]:
            p(f"    {s['drop']:+.0%}  價 {s['px']:7.2f}  加碼 {s['amt']:,.0f}"
              f"  累計 {s['cum']:,.0f}  約 {s['shares']:,} 股")
        ok = "足夠" if r["cash"] >= r["ladder_total"] else "不足"
        p(f"    需求合計 {r['ladder_total']:,.0f} vs 現金 {r['cash']:,.0f}:{ok}")


def holdings_from_wb(wb, tab="03_持股總表"):
    ws = wb[tab]
    out = []
    for r in range(3, ws.max_row + 1):
        n = ws.cell(r, 1).value
        if not n or str(n).strip() in ("", "合計"):
            continue
        out.append({"name": str(n).strip(), "cat": str(ws.cell(r, 2).value).strip(),
                    "sh": float(ws.cell(r, 4).value), "px": float(ws.cell(r, 12).value)})
    return out


def main():
    import argparse
    from datetime import datetime
    import openpyxl
    ap = argparse.ArgumentParser(description="對本機 Wealth OS xlsm 計算正二資產負債表")
    ap.add_argument("--xlsm", required=True)
    ap.add_argument("--date", default=f"{datetime.now():%Y/%m/%d}")
    args = ap.parse_args()
    wb = openpyxl.load_workbook(args.xlsm, keep_vba=True)
    try:
        cfg = read_lev2_cfg(wb)
    except ValueError as e:
        sys.exit(f"參數不完整:{e}")
    print_report(compute(holdings_from_wb(wb), cfg, prev_month_mv(wb, args.date)))


if __name__ == "__main__":
    main()
